"""Runs without an API key: validation rules, id/RTM/stats derivation, and both exports.
    python -m pytest backend/tests   or   python backend/tests/test_pipeline.py
"""

from __future__ import annotations

import csv
import io
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1]))

from openpyxl import load_workbook  # noqa: E402

from app.export import COLUMNS, to_csv, to_xlsx  # noqa: E402
from app.generate import ApiError, finalize, tidy, validate  # noqa: E402
from app.jira import adf_to_text, enabled, parse_key  # noqa: E402
from app.schemas import GenerateRequest, LlmOutput  # noqa: E402

FIXTURE = json.loads((Path(__file__).parent / "fixture_llm_output.json").read_text(encoding="utf-8"))
REQ = GenerateRequest(
    requirement_text="Tambahkan tombol 'Cek Limit dan Akad SIKP' pada Step 1 halaman Initial Application Terms untuk pengajuan produk KUR.",
    system_under_test="PEGA", story_key="DEMO-5385", story_title="Pega Micro - Mapping Bunga Graduasi",
    story_url="https://example.atlassian.net/browse/DEMO-5385", language="id", req_prefix="TC-03",
)


def result():
    out = LlmOutput.model_validate(FIXTURE)
    tidy(out.groups)
    errors, warnings = validate(out.groups, "id")
    assert errors == [], errors
    return finalize(out, REQ, "test-model", warnings)


def test_finalize_ids_order_and_stats():
    r = result()
    g1, g2 = r.groups
    assert (g1.req_no, g2.req_no) == ("TC-03", "TC-04")  # req_prefix drives the first group number
    assert [c.tc_id for c in g1.test_cases] == ["TC-03-01", "TC-03-02", "TC-03-03", "TC-03-04"]
    assert [c.case_class for c in g1.test_cases] == ["positive", "positive", "negative", "boundary"]  # positives first, stable
    assert [c.pos_neg for c in g1.test_cases] == ["Positive", "Positive", "Negative", "Negative"]  # boundary exports as Negative
    assert all(c.status == "" and c.tested_by == "" and c.execution_date is None and c.notes == "" for g in r.groups for c in g.test_cases)
    assert r.stats.model_dump() == {
        "total_cases": 6, "by_class": {"positive": 3, "negative": 2, "boundary": 1},
        "by_priority": {"High": 5, "Medium": 1, "Low": 0}, "ac_total": 6, "ac_covered": 5, "ac_uncovered": ["AC-6"],
    }
    assert "AC-6 produced no test case" in r.warnings
    assert r.rtm[0].model_dump() == {"req_no": "TC-03", "requirement": "DEMO-5385 | Group 1: Tombol 'Cek Limit dan Akad SIKP' dan Validasi Tombol 'Hitung Angsuran'",
                                    "total_case": 4, "passed": 0, "failed": 0, "not_tested": 4, "completeness": 0.0}
    assert r.notes[0].kind == "source" and r.notes[0].label == "DEMO-5385"


def test_validation_rules():
    out = LlmOutput.model_validate(FIXTURE)
    tidy(out.groups)
    c = out.groups[0].test_cases[0]
    c.title = "Tombol tampil"                       # positive without prefix
    c.expected_result = "Berhasil."                 # bare denylist word
    c.test_steps = ["Cek tombol"]                   # no entry point
    c.requirement_ref = "AC-1, AC-2"                # tidy() keeps only the first id, so this one passes
    tidy(out.groups)
    assert c.requirement_ref == "AC-1"
    errors, _ = validate(out.groups, "id")
    assert len(errors) == 3, errors
    assert any("must start with 'Memastikan'" in e for e in errors)
    assert any("not verifiable" in e for e in errors)
    assert any("login/entry point" in e for e in errors)
    # English prefix rule
    out = LlmOutput.model_validate(FIXTURE)
    out.groups[0].test_cases = out.groups[0].test_cases[:1]
    out.groups[0].test_cases[0].title = "Verify the button is displayed"
    errors, _ = validate(out.groups[:1], "en")
    assert not any("must start with" in e for e in errors)
    # all-positive group is a warning, not an error
    out = LlmOutput.model_validate(FIXTURE)
    out.groups[1].test_cases = out.groups[1].test_cases[:1]
    errors, warnings = validate(out.groups, "id")
    assert errors == [] and any("only positive cases" in w for w in warnings)


def test_exports():
    r = result()
    rows = list(csv.reader(io.StringIO(to_csv(r).lstrip("﻿"))))
    assert rows[0] == COLUMNS
    assert rows[1][0] == "DEMO-5385 | Pega Micro - Mapping Bunga Graduasi"
    assert rows[2][0].startswith("-- Group 1:")
    assert rows[3][0] == "TC-03-01" and rows[3][6].startswith("1. Login PEGA") and "\n2. " in rows[3][6]
    assert len(rows) == 1 + 1 + 2 + 6  # header, story, 2 groups, 6 cases

    wb = load_workbook(io.BytesIO(to_xlsx(r, "Proj", "No-1", "Me")))
    assert wb.sheetnames == ["RTM", "Test Script", "Defect List", "Note"]
    ts = wb["Test Script"]
    assert [c.value for c in ts[1]] == COLUMNS
    assert ts["A2"].value.startswith("DEMO-5385 |") and ts["A3"].value.startswith("-- Group 1:") and ts["A4"].value == "TC-03-01"
    assert ts["I4"].value is None and ts.freeze_panes == "A2"
    rtm = wb["RTM"]
    assert rtm["B1"].value == "Proj" and rtm["A24"].value == "Req No"
    assert rtm["A25"].value == "TC-03" and rtm["C25"].value == "=COUNTIFS('Test Script'!A:A, \"TC-03*\")"
    assert rtm["A28"].value == "TC-04" and rtm["B10"].value == "=SUM(C25:C30)"
    assert wb["Defect List"]["A17"].value == "No"
    note = wb["Note"]
    assert note["A1"].value == "Sumber Story" and note["A2"].value == "DEMO-5385"
    assert any(note.cell(i, 1).value == "Asumsi yang perlu dikonfirmasi" for i in range(1, 15))


def test_jira_parse_key():
    assert parse_key("DEMO-5385") == "DEMO-5385"
    assert parse_key("  demo-5385 ") == "DEMO-5385"
    assert parse_key("https://your-company.atlassian.net/browse/DEMO-5385") == "DEMO-5385"
    assert parse_key("https://x.atlassian.net/browse/MB-22?filter=1") == "MB-22"
    for bad in ("", "not a ticket", "12345", "https://example.com/"):
        try:
            parse_key(bad)
            raise AssertionError(f"{bad!r} should have been rejected")
        except ApiError as e:
            assert e.code == "jira_bad_key"


def test_jira_adf_to_text():
    doc = json.loads((Path(__file__).parent / "fixture_adf.json").read_text(encoding="utf-8"))
    text, media = adf_to_text(doc)
    assert media == 3  # mediaSingle + two inside mediaGroup
    assert "### Accaptance Criteria :" in text
    assert "Tombol 'Cek Limit dan akad' ditambahkan di Step 1." in text      # inline marks flattened, spacing kept
    assert "Kolom | Keterangan" in text                                      # table header row survives
    assert "Kel_Sektor | 1 = Produksi, 2 = Perdagangan Non Ekspor" in text   # table body survives
    assert "- Akad ke 1 -> bunga efektif = 6%" in text                       # list marker
    assert "- subsidi bunga = 4,5%" in text                                  # nested list still reaches the text
    assert "Teks di dalam node tak dikenal" in text                          # unknown node recurses, does not vanish
    assert "https://example.com/spec" in text                                # inlineCard becomes its url
    assert adf_to_text(None) == ("", 0) and adf_to_text({}) == ("", 0)       # empty description is not a crash


def test_jira_disabled_without_credentials():
    for k in ("JIRA_BASE_URL", "JIRA_EMAIL", "JIRA_API_TOKEN"):
        os.environ.pop(k, None)
    assert enabled() is False                                                # public demo default: feature off
    os.environ.update(JIRA_BASE_URL="https://x.atlassian.net", JIRA_EMAIL="a@b.c")
    assert enabled() is False                                                # partial config is still off
    os.environ["JIRA_API_TOKEN"] = "t"
    assert enabled() is True
    for k in ("JIRA_BASE_URL", "JIRA_EMAIL", "JIRA_API_TOKEN"):
        os.environ.pop(k, None)


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
