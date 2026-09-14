"""XLSX (RTM, Test Script, Defect List, Note) and CSV export.
Layout and formatting copied from the real bank SIT workbook in excel_format/."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .schemas import GenerationResult

COLUMNS = ["TC ID", "Title / Description", "Pos/Neg", "Priority", "Type", "Precondition",
           "Test Steps", "Expected Result", "Status", "Tested By", "Execution Date", "Notes"]
WIDTHS = [12.4, 39.1, 11, 9.9, 11.6, 34, 43.4, 40.7, 24.6, 11, 16.3, 14.1]

ARIAL = Font(name="Arial", size=11)
BOLD = Font(name="Arial", size=11, bold=True)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, horizontal="center", vertical="center")
LEFT = Alignment(wrap_text=True, horizontal="left", vertical="center")


def fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb)


HEADER = fill("A4C2F4")
STORY = fill("F4CCCC")
GROUP = fill("FFF2CC")
BLUE, GREEN, RED, GREY, YELLOW, DARK, LABEL = (fill(c) for c in ("C9DAF8", "D9EAD3", "F4CCCC", "EFEFEF", "FFF2CC", "BDBDBD", "CFE2F3"))


def steps_text(steps: list[str]) -> str:
    return "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1))


def script_rows(result: GenerationResult) -> list[tuple[str, list]]:
    """Row stream shared by CSV and XLSX: ('story'|'group'|'case', values)."""
    rows: list[tuple[str, list]] = []
    story = " | ".join(x for x in (result.story.story_key, result.story.story_title) if x)
    if story:
        rows.append(("story", [story]))
    for g in result.groups:
        rows.append(("group", [f"-- Group {g.group_no}: {g.group_label} --"]))
        for c in g.test_cases:
            rows.append(("case", [c.tc_id, c.title, c.pos_neg, c.priority, c.type, c.precondition,
                                  steps_text(c.test_steps), c.expected_result, c.status, c.tested_by,
                                  c.execution_date or "", c.notes]))
    return rows


def to_csv(result: GenerationResult) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(COLUMNS)
    for _, vals in script_rows(result):
        w.writerow(vals + [""] * (len(COLUMNS) - len(vals)))
    return "﻿" + buf.getvalue()  # BOM so Excel opens UTF-8 correctly


def to_xlsx(result: GenerationResult, project_name: str = "", project_no: str = "", created_by: str = "") -> bytes:
    wb = Workbook()
    ws_rtm = wb.active
    ws_rtm.title = "RTM"
    _test_script(wb.create_sheet("Test Script"), result)
    _rtm(ws_rtm, result, project_name, project_no, created_by)
    _defect_list(wb.create_sheet("Defect List"))
    _note(wb.create_sheet("Note"), result)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _put(ws: Worksheet, ref: str, value, font=BOLD, fill_=None, align=CENTER, number_format=None):
    c = ws[ref]
    c.value = value
    c.font, c.alignment, c.border = font, align, BORDER
    if fill_:
        c.fill = fill_
    if number_format:
        c.number_format = number_format
    return c


def _test_script(ws: Worksheet, result: GenerationResult) -> None:
    for i, (h, w) in enumerate(zip(COLUMNS, WIDTHS), 1):
        c = ws.cell(1, i, h)
        c.font, c.fill, c.alignment, c.border = BOLD, HEADER, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    r = 2
    for kind, vals in script_rows(result):
        if kind == "case":
            for col, v in enumerate(vals, 1):
                c = ws.cell(r, col, v)
                c.font, c.alignment, c.border = (BOLD if col == 1 else ARIAL), TOP, BORDER
        else:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLUMNS))
            c = ws.cell(r, 1, vals[0])
            c.font, c.fill, c.alignment = BOLD, (STORY if kind == "story" else GROUP), TOP
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(r, col).border = BORDER
        r += 1
    dv = DataValidation(type="list", formula1='"Not Tested,Passed,Failed,On Hold"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{max(r - 1, 2)}")


def _rtm(ws: Worksheet, result: GenerationResult, project_name: str, project_no: str, created_by: str) -> None:
    for col, w in zip("ABCDEFG", (24.1, 85.7, 12, 10, 10, 12, 15.4)):
        ws.column_dimensions[col].width = w
    small = Font(name="Arial", size=10, bold=True)
    for row, label, value in ((1, "Project Name", project_name), (3, "Project No", project_no), (5, "Created By", created_by)):
        _put(ws, f"A{row}", label, small, LABEL)
        _put(ws, f"B{row}", value, small, fill("FFFFFF"))

    first, last = 25, 25 + 3 * max(len(result.rtm), 1) - 1  # 3 merged rows per RTM entry, like the template
    for ref, v in (("A9", "Status"), ("B9", "Total"), ("C9", "Percentage")):
        _put(ws, ref, v, fill_=DARK)
    summary = (
        (10, "Total Case", f"=SUM(C{first}:C{last})", BLUE, None),
        (12, "Passed", f"=SUM(D{first}:D{last})", GREEN, "=IFERROR(B12/B10,0)"),
        (14, "Failed", f"=SUM(E{first}:E{last})", RED, "=IFERROR(B14/B10,0)"),
        (16, "Not Tested", f"=SUM(F{first}:F{last})", GREY, "=IFERROR(B16/B10,0)"),
        (18, "Defect", "=COUNTA('Defect List'!A18:A120)", YELLOW, "=IFERROR(B18/B10,0)"),
    )
    for row, label, formula, f_, pct in summary:
        _put(ws, f"A{row}", label, fill_=f_)
        _put(ws, f"B{row}", formula, fill_=f_)
        if pct:
            _put(ws, f"C{row}", pct, number_format="0.0%")
        for col in "ABC":
            ws.merge_cells(f"{col}{row}:{col}{row + 1}")

    _put(ws, "A22", "Requirement Traceability Matrix", Font(name="Arial", size=15, bold=True), BLUE)
    for col, h in zip("ABCDEFG", ("Req No", "Requirement", "Total Case", "Passed", "Failed", "Not Tested", "Completeness")):
        _put(ws, f"{col}24", h)

    r = first
    for row in result.rtm:
        q = f"'Test Script'!A:A, \"{row.req_no}*\""
        cells = (
            ("A", row.req_no, None, CENTER),
            ("B", row.requirement, None, LEFT),
            ("C", f"=COUNTIFS({q})", BLUE, CENTER),
            ("D", f"=COUNTIFS({q}, 'Test Script'!I:I, \"Passed\")", GREEN, CENTER),
            ("E", f"=COUNTIFS({q}, 'Test Script'!I:I, \"Failed\")", RED, CENTER),
            ("F", f"=COUNTIFS({q}, 'Test Script'!I:I, \"Not Tested\")", fill("F3F3F3"), CENTER),
            ("G", f'=IF(C{r}=0, "0%", REPT("█", ROUND((D{r}/C{r})*10,0)) & " " & TEXT(D{r}/C{r}, "0%") & " " & REPT("░", 10-ROUND((D{r}/C{r})*10,0)))', None, CENTER),
        )
        for col, v, f_, align in cells:
            _put(ws, f"{col}{r}", v, fill_=f_, align=align)
            ws.merge_cells(f"{col}{r}:{col}{r + 2}")
        r += 3


def _defect_list(ws: Worksheet) -> None:
    for col, w in zip("ABCDEFGH", (21.6, 51.9, 37.1, 18, 20.6, 18, 16, 37.4)):
        ws.column_dimensions[col].width = w
    for row, label, formula in ((3, "Total Defect", "=COUNTA(G18:G1000)"), (5, "Done", '=COUNTIF(G18:G1000,"Done")'),
                                (7, "In Progress", '=COUNTIF(G18:G1000,"In Progress")'), (9, "Open", '=COUNTIF(G18:G1000,"Open")'),
                                (11, "Dropped", '=COUNTIF(G18:G1000,"Dropped")')):
        ws[f"A{row}"].value, ws[f"A{row}"].font = label, BOLD
        ws[f"B{row}"].value, ws[f"B{row}"].font = formula, BOLD
    headers = ["No", "Description", "Defect ticket Number", "Raised Date", "Resolved Date",
               "TTR (Time to Resolve) in Hours", "Status", "Note SIT"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(17, i, h)
        c.font, c.fill, c.alignment, c.border = BOLD, HEADER, CENTER, BORDER


NOTE_SECTIONS = (
    ("source", {"id": "Sumber Story", "en": "Source Story"}),
    ("correction", {"id": "Koreksi terhadap teks AC", "en": "Corrections against the AC text"}),
    ("assumption", {"id": "Asumsi yang perlu dikonfirmasi", "en": "Assumptions to confirm"}),
    ("dedup", {"id": "Deduplikasi", "en": "Deduplication"}),
)


def _note(ws: Worksheet, result: GenerationResult) -> None:
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110
    r = 1
    for kind, title in NOTE_SECTIONS:
        items = [n for n in result.notes if n.kind == kind]
        if not items:
            continue
        ws.cell(r, 1, title[result.language]).font = BOLD
        r += 1
        for n in items:
            ws.cell(r, 1, n.label).font = ARIAL
            c = ws.cell(r, 2, n.detail)
            c.font, c.alignment = ARIAL, Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1
